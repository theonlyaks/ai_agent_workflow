"""Add (or remove) the demo 11th workflow in the real Excel file + YAML bindings.

    python scripts/add_wf011_demo.py            # add WF011
    python scripts/add_wf011_demo.py --remove   # remove WF011
"""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook
import yaml

ROOT = Path(__file__).resolve().parents[1]
EXCEL = ROOT / "data" / "AI_Agent_Workflow_Assessment.xlsx"
CONFIG = ROOT / "config" / "workflows.yaml"
EXAMPLE = ROOT / "examples" / "wf011_low_margin"
MARKER = "  # --- WF011 demo (added by scripts/add_wf011_demo.py) ---\n"


def add() -> None:
    row = yaml.safe_load((EXAMPLE / "excel_row.yaml").read_text(encoding="utf-8"))
    workbook = load_workbook(EXCEL)
    sheet = workbook["Workflows"]
    headers = [cell.value for cell in sheet[1]]
    if any(r[0] == row["Workflow_ID"] for r in sheet.iter_rows(min_row=2, values_only=True)):
        print("WF011 already in Excel")
    else:
        sheet.append([row.get(h, "") for h in headers])  # openpyxl keeps the other sheets/formatting
        workbook.save(EXCEL)
        print("Added WF011 row to", EXCEL.name)

    text = CONFIG.read_text(encoding="utf-8")
    if MARKER in text:
        print("WF011 binding already in", CONFIG.name)
        return
    binding = (EXAMPLE / "binding.yaml").read_text(encoding="utf-8")
    body = "".join("  " + line if line.strip() and not line.startswith("#") else line
                   for line in binding.splitlines(keepends=True) if not line.startswith("#"))
    CONFIG.write_text(text.rstrip("\n") + "\n\n" + MARKER + body, encoding="utf-8")
    print("Added WF011 binding to", CONFIG.name)


def remove() -> None:
    workbook = load_workbook(EXCEL)
    sheet = workbook["Workflows"]
    for index in range(sheet.max_row, 1, -1):
        if sheet.cell(row=index, column=1).value == "WF011":
            sheet.delete_rows(index)
    workbook.save(EXCEL)
    text = CONFIG.read_text(encoding="utf-8")
    if MARKER in text:
        CONFIG.write_text(text.split("\n\n" + MARKER)[0].rstrip("\n") + "\n", encoding="utf-8")
    print("Removed WF011 from", EXCEL.name, "and", CONFIG.name)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--remove", action="store_true")
    remove() if parser.parse_args().remove else add()
